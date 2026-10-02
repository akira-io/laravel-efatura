"""Build runtime JSON catalogs from the bundled, immutable DNRE artifacts.

Requires Python 3.12, openpyxl 3.1.5, and xlrd 2.0.2 at build time only.
Run with --check to verify committed JSON without changing it.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl
import xlrd


ROOT = Path(__file__).resolve().parents[1]
RESOURCES = ROOT / "resources"
SOURCE = RESOURCES / "catalogs" / "source"
XSD = RESOURCES / "xsd" / "efatura" / "2024-05-27" / "common"
NS = {"x": "http://www.w3.org/2001/XMLSchema", "c": "urn:un:unece:uncefact:documentation:standard:CoreComponentsTechnicalSpecification:2"}
ARCHIVE = SOURCE / "2024-05-27-XML-XSD.zip"


def source(path: Path) -> dict[str, str]:
    return {
        "path": path.relative_to(RESOURCES).as_posix(),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def clean(value: object) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (str, int, float)):
        return value
    raise TypeError(f"Unexpected worksheet value: {value!r}")


def xsd_rows(filename: str) -> list[dict[str, str]]:
    tree = ET.parse(XSD / filename)
    rows = []
    for entry in tree.findall(".//x:enumeration", NS):
        row = {"code": entry.attrib["value"]}
        name = entry.find(".//c:Name", NS)
        if name is not None and name.text:
            row["name"] = name.text.strip()
        rows.append(row)
    return rows


def units() -> tuple[list[dict], dict]:
    path = SOURCE / "codigos-de-unidades-de-medidas.xls"
    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_name("Unidades Completas")
    fields = ("status", "code", "name", "description", "level_category", "symbol", "conversion_factor")
    records = [dict(zip(fields, (clean(v) for v in sheet.row_values(i)), strict=True)) for i in range(1, sheet.nrows)]
    common = workbook.sheet_by_name("Unidades Mais Usadas")
    common_fields = ("category", "code", "name", "symbol", "notes")
    common_rows = [dict(zip(common_fields, (clean(v) for v in common.row_values(i)), strict=True)) for i in range(1, common.nrows)]
    levels = workbook.sheet_by_name("Níveis-Categorias")
    level_rows = [{"level_category": clean(levels.cell_value(i, 0)), "description": clean(levels.cell_value(i, 1))} for i in range(1, levels.nrows)]
    return records, {"sources": [source(path)], "common_rows": common_rows, "level_rows": level_rows}


def places() -> tuple[list[dict], list[dict], Path]:
    path = SOURCE / "codigo-paises-lugares-cv.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["CODIGOS"]
    rows = sheet.iter_rows(values_only=True)
    fields = tuple(str(v).lower() for v in next(rows))
    records = [dict(zip(fields, (clean(v) for v in row), strict=True)) for row in rows]
    countries = [row for row in records if row["nivel"] == 1]
    workbook.close()
    return records, countries, path


def tax_exemptions() -> tuple[list[dict], Path]:
    path = SOURCE / "Lista-de-Motivos-de-Nao-Liquidacao-de-Imposto.xlsx"
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Motivos"]
    rows = list(sheet.iter_rows(values_only=True))[1:]
    records = [{"code": str(row[0]), "description": clean(row[1]), "mention": clean(row[2])} for row in rows if row[0] is not None]
    workbook.close()
    return records, path


def build() -> dict[str, dict]:
    unit_rows, unit_metadata = units()
    place_rows, published_countries, place_path = places()
    country_names = {str(row["codigo"]): str(row["nome"]) for row in published_countries}
    country_xsd = "ISO_ISOTwo-letterCountryCode_SecondEdition2006.xsd"
    country_rows = xsd_rows(country_xsd)
    for row in country_rows:
        if row["code"] in country_names:
            row["published_name"] = country_names[row["code"]]
    exemption_rows, exemption_path = tax_exemptions()
    exemption_xsd = "CV_EFatura_TaxExemptionReason_v1.0.xsd"
    assert [row["code"] for row in exemption_rows] == [row["code"] for row in xsd_rows(exemption_xsd)]
    currency_xsd = "ISO_ISO3AlphaCurrencyCode_2012-08-31.xsd"
    payment_xsd = "UNECE_PaymentMeansCode_D19B.xsd"
    return {
        "units": {"records": unit_rows, **unit_metadata},
        "countries": {"records": country_rows, "sources": [source(XSD / country_xsd), source(place_path), source(ARCHIVE)], "published_country_rows": published_countries},
        "locations": {"records": place_rows, "sources": [source(place_path)], "accepted_count": sum(row["nivel"] > 1 and str(row["codigo"]).startswith("CV") for row in place_rows)},
        "currencies": {"records": xsd_rows(currency_xsd), "sources": [source(XSD / currency_xsd), source(ARCHIVE)]},
        "payment_means": {"records": xsd_rows(payment_xsd), "sources": [source(XSD / payment_xsd), source(ARCHIVE)]},
        "tax_exemption_reasons": {"records": exemption_rows, "sources": [source(exemption_path), source(XSD / exemption_xsd), source(ARCHIVE)]},
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Compare generated bytes with committed JSON")
    args = parser.parse_args()
    for name, payload in build().items():
        records = payload["records"]
        codes = [record.get("code", record.get("codigo")) for record in records]
        assert len(codes) == len(set(codes)), f"Duplicate codes in {name}"
        path = RESOURCES / "catalogs" / f"{name}.json"
        content = (json.dumps({"schema_version": 1, "count": len(records), **payload}, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()
        if args.check:
            assert path.read_bytes() == content, f"Catalog differs from sources: {name}"
        else:
            path.write_bytes(content)
        print(f"{name}: {len(records)} records, sha256 {hashlib.sha256(content).hexdigest()}")


if __name__ == "__main__":
    main()
